import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

MATH_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

def set_page_margins(doc, top_cm=1.0, bottom_cm=1.0, left_cm=1.0, right_cm=1.0):
    for sec in doc.sections:
        sec.top_margin = Inches(top_cm / 2.54)
        sec.bottom_margin = Inches(bottom_cm / 2.54)
        sec.left_margin = Inches(left_cm / 2.54)
        sec.right_margin = Inches(right_cm / 2.54)

def setup_footer_page_number(doc):
    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        r1 = p.add_run("〔第")
        set_run_font(r1, font_name="標楷體", font_size=10)
        fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        p._p.append(fld1)
        r2 = p.add_run("頁,共")
        set_run_font(r2, font_name="標楷體", font_size=10)
        fld2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        p._p.append(fld2)
        r3 = p.add_run("頁〕")
        set_run_font(r3, font_name="標楷體", font_size=10)

def set_run_font(run, font_name="標楷體", ascii_font="Times New Roman", font_size=11, bold=False, italic=False, underline=False, color=None):
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{ascii_font}" w:hAnsi="{ascii_font}" w:eastAsia="{font_name}"/>')
    rPr.append(rFonts)

def m_run(text, italic=True):
    sty = '<m:rPr><m:sty m:val="p"/></m:rPr>' if not italic else ''
    esc = str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return f'<m:r>{sty}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math" w:eastAsia="標楷體"/><w:sz w:val="22"/></w:rPr><m:t>{esc}</m:t></m:r>'

def m_frac(num_xml, den_xml):
    return f'<m:f><m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>'

def m_sqrt(base_xml, deg_xml=None):
    if deg_xml is None:
        return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{base_xml}</m:e></m:rad>'
    return f'<m:rad><m:radPr><m:degHide m:val="0"/></m:radPr><m:deg>{deg_xml}</m:deg><m:e>{base_xml}</m:e></m:rad>'

def m_sup(base_xml, sup_xml):
    return f'<m:sSup><m:e>{base_xml}</m:e><m:sup>{sup_xml}</m:sup></m:sSup>'

def m_bar(base_xml, pos="top"):
    return f'<m:bar><m:barPr><m:pos m:val="{pos}"/></m:barPr><m:e>{base_xml}</m:e></m:bar>'

def append_omath(paragraph, inner_omml_xml):
    omath_str = f'<m:oMath {MATH_NS}>{inner_omml_xml}</m:oMath>'
    paragraph._p.append(parse_xml(omath_str))

def add_jinhe_header(doc, doc_subtitle="第三次段考試題"):
    # 1. 抬頭：以粗體加底線打上〔新北市立錦和高級中學115學年度第一學期國中部○年級○○科第三次段考試題〕
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.left_indent = Pt(0)
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(3)
    title_p.paragraph_format.line_spacing = 1.15
    r_title = title_p.add_run(f"新北市立錦和高級中學115學年度第一學期國中部九年級數學科{doc_subtitle}")
    set_run_font(r_title, font_name="標楷體", font_size=13.5, bold=True, underline=True)

    # 2. 命題範圍與學生資訊
    info_p = doc.add_paragraph()
    info_p.paragraph_format.left_indent = Pt(0)
    info_p.paragraph_format.space_before = Pt(0)
    info_p.paragraph_format.space_after = Pt(3)
    info_p.paragraph_format.line_spacing = 1.15
    r_range = info_p.add_run("﹝命題範圍：九上 第 3 章 幾何與證明（3-1 證明與推理、3-2 三角形的外心、內心與重心）﹞")
    set_run_font(r_range, font_name="標楷體", font_size=10.5, bold=True)
    r_space = info_p.add_run("   ")
    r_stu = info_p.add_run("班級：________  座號：____  姓名：____________")
    set_run_font(r_stu, font_name="標楷體", font_size=10.5, bold=True)

    # 3. 規範 5：答案卡警語（紅字粗體）
    w1_p = doc.add_paragraph()
    w1_p.paragraph_format.left_indent = Pt(0)
    w1_p.paragraph_format.space_before = Pt(0)
    w1_p.paragraph_format.space_after = Pt(2)
    w1_p.paragraph_format.line_spacing = 1.15
    r_w1 = w1_p.add_run("※ 答案卷(卡)未寫班級、姓名、座號，或畫卡錯誤致電腦無法判讀考生身份者，一律扣5分。")
    set_run_font(r_w1, font_name="標楷體", font_size=10, bold=True, color=RGBColor(200, 0, 0))

    # 4. 規範 7：若有出選擇題與非選擇題於同一張答案卷時
    w2_p = doc.add_paragraph()
    w2_p.paragraph_format.left_indent = Pt(0)
    w2_p.paragraph_format.space_before = Pt(0)
    w2_p.paragraph_format.space_after = Pt(4)
    w2_p.paragraph_format.line_spacing = 1.15
    r_w2 = w2_p.add_run("※ 本試卷所有試題（選擇題、填充題、非選擇題）請一律在答案卷上作答，請一律用黑色墨水筆作答，違者扣總分5分。")
    set_run_font(r_w2, font_name="標楷體", font_size=10, bold=True)

def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, font_name="標楷體", font_size=11.5, bold=True)
    return p

def create_question_p(doc, num_str, text_prefix, scope_str=None, space_after=2):
    """
    建立題號凸排 (Hanging Indent) 之題幹段落：
    - left_indent = Inches(0.28)
    - first_line_indent = -Inches(0.28)
    - 右對齊定位點 (7.48 in)：出題範圍 (3-X) 取消粗體 (bold=False)
    """
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.48), WD_TAB_ALIGNMENT.RIGHT)

    r_num = p.add_run(num_str)
    set_run_font(r_num, font_name="Times New Roman", font_size=11, bold=True)
    r_txt = p.add_run(text_prefix)
    set_run_font(r_txt, font_name="標楷體", font_size=11)

    if scope_str:
        r_tab = p.add_run(f"\t{scope_str}")
        # 出題範圍字型取消加粗體 (bold=False)
        set_run_font(r_tab, font_name="Times New Roman", font_size=10.5, bold=False)

    return p

def create_options_p(doc, opt_text, space_after=3):
    """
    建立選項段落，配合題號凸排，左側縮排 Inches(0.28)。
    """
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(opt_text)
    set_run_font(r, font_name="Times New Roman", font_size=11)
    return p

def add_question_with_image_top_bottom(doc, num_str, text_prefix, scope_str, image_path, options_str=None, img_width_in=2.1, space_after=3):
    """
    依「上及下」文繞圖規範插入附圖題目：
    1. 題幹在上方（題號凸排，(3-X) 靠右無加粗）
    2. 圖形居中（圖片與圖題具備充裕間距，文字無重疊）
    3. 選項在下方（左縮排對齊選項）
    """
    # 1. 題幹段落
    p_q = create_question_p(doc, num_str, text_prefix, scope_str, space_after=2)

    # 2. 圖形段落 (置中，上下型排版)
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.left_indent = Pt(0)
        p_img.paragraph_format.first_line_indent = Pt(0)
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(image_path, width=Inches(img_width_in))

    # 3. 選項段落 (若有)
    if options_str:
        if isinstance(options_str, list):
            for opt_line in options_str:
                create_options_p(doc, opt_line, space_after=2)
        else:
            create_options_p(doc, options_str, space_after=space_after)

def safe_save(doc, filename):
    try:
        doc.save(filename)
        print(f"[OK] 檔案已成功生成：{filename}")
        return filename
    except PermissionError:
        base, ext = os.path.splitext(filename)
        alt_filename = f"{base}_修訂版{ext}"
        doc.save(alt_filename)
        print(f"[NOTE] 因 {filename} 正在被 Word 開啟鎖定，已儲存為：{alt_filename}")
        return alt_filename

print('Updated exam_helpers.py with new Jinhe standards ready!')
